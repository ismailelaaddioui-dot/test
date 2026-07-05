"""Extract raw text from local files (PDF, DOCX, CSV, TXT).

Each parser is defensive: unreadable or corrupt files raise
`DocumentParseError` instead of propagating library-specific exceptions,
so callers can log-and-continue rather than crash.
"""

import csv
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv", ".txt"}


class DocumentParseError(Exception):
    """Raised when a local file can't be read or parsed."""


def parse_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:  # ImportError, or pypdf's own import-time failures
        raise DocumentParseError(f"pypdf is required to read PDF files: {exc}") from exc

    try:
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # pypdf can raise many different error types
        raise DocumentParseError(f"failed to parse PDF: {exc}") from exc


def parse_docx(path: Path) -> str:
    try:
        import docx
    except Exception as exc:  # ImportError, or python-docx's own import-time failures
        raise DocumentParseError(f"python-docx is required to read DOCX files: {exc}") from exc

    try:
        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells)
        return "\n".join(paragraphs)
    except Exception as exc:
        raise DocumentParseError(f"failed to parse DOCX: {exc}") from exc


def parse_csv(path: Path) -> str:
    try:
        with open(path, newline="", encoding="utf-8", errors="replace") as fh:
            reader = csv.reader(fh)
            return "\n".join(", ".join(row) for row in reader)
    except OSError as exc:
        raise DocumentParseError(f"failed to read CSV: {exc}") from exc


def parse_txt(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise DocumentParseError(f"failed to read TXT: {exc}") from exc


_PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".csv": parse_csv,
    ".txt": parse_txt,
}


def parse_file(path: Path) -> str:
    """Dispatch to the right parser based on file extension."""
    suffix = path.suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise DocumentParseError(f"unsupported file type: {suffix}")
    return parser(path)
